"""Build the report from Markdown using the managed document runtime.

Equations are native OMML. Tables and figures stay editable/in their source form.
"""
from __future__ import annotations

import copy
import json
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from PIL import Image

DIR = Path(__file__).resolve().parent
STEM = '圆柱绕流符号建模与灵敏度研究成果汇总报告_2026-09-20'
MD = DIR / (STEM+'.md')
OUT = DIR / (STEM+'.docx')
ASSETS = DIR / '成果汇总附件_2026-09-20'
WIDTH_CM = 16.6


def font_style(st, size, bold=False, chinese='宋体'):
    st.font.name = 'Calibri'
    st.font.size = Pt(size)
    st.font.bold = bold
    st.font.color.rgb = RGBColor(0,0,0)
    pr=st.element.get_or_add_rPr()
    rf=pr.find(qn('w:rFonts'))
    if rf is None:
        rf=OxmlElement('w:rFonts');pr.append(rf)
    rf.set(qn('w:eastAsia'),chinese)
    for a in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
        rf.attrib.pop(qn('w:'+a),None)
    lang=OxmlElement('w:lang');lang.set(qn('w:val'),'en-US');lang.set(qn('w:eastAsia'),'zh-CN');pr.append(lang)


def inline(p,text):
    # Relative Markdown links retain their human-readable labels in Word.
    text=re.sub(r'\[([^\]]+)\]\(([^)]+)\)',r'\1',text)
    parts=re.split(r'(\*\*.*?\*\*|`[^`]+`)',text)
    for part in parts:
        if not part:continue
        bold=part.startswith('**') and part.endswith('**')
        mono=part.startswith('`') and part.endswith('`')
        if bold:part=part[2:-2]
        if mono:part=part[1:-1]
        # Give long archived filenames legitimate soft wrap opportunities.
        if len(part)>35:
            part=re.sub(r'(?<=[/_])(?=[A-Za-z0-9])','\u200b',part)
        rr=p.add_run(part);rr.bold=bold
        if mono:rr.font.name='Consolas';rr.font.size=Pt(9)


def atom(text):
    r=OxmlElement('m:r')
    pr=OxmlElement('w:rPr')
    ff=OxmlElement('w:rFonts');ff.set(qn('w:ascii'),'Cambria Math');ff.set(qn('w:hAnsi'),'Cambria Math');pr.append(ff)
    sz=OxmlElement('w:sz');sz.set(qn('w:val'),'22');pr.append(sz)
    r.append(pr)
    t=OxmlElement('m:t');t.text=text
    if text.startswith(' ') or text.endswith(' '):t.set(qn('xml:space'),'preserve')
    r.append(t);return r


def seq(*args):
    rr=[]
    for a in args:
        if isinstance(a,str):rr.append(atom(a))
        elif isinstance(a,list):rr.extend(copy.deepcopy(a))
        else:rr.append(copy.deepcopy(a))
    return rr


def item(tag,children):
    x=OxmlElement('m:'+tag)
    for child in seq(children):x.append(child)
    return x


def scripts(base,sub=None,sup=None):
    tag='sSubSup' if sub is not None and sup is not None else ('sSub' if sub is not None else 'sSup')
    x=OxmlElement('m:'+tag);x.append(item('e',base))
    if sub is not None:x.append(item('sub',sub))
    if sup is not None:x.append(item('sup',sup))
    return x


def frac(top,bottom):
    f=OxmlElement('m:f');f.append(item('num',top));f.append(item('den',bottom));return f


def paren(content):
    d=OxmlElement('m:d');pr=OxmlElement('m:dPr')
    for tag,ch in [('begChr','('),('endChr',')')]:
        e=OxmlElement('m:'+tag);e.set(qn('m:val'),ch);pr.append(e)
    d.append(pr);d.append(item('e',content));return d


def qs():return scripts('q',sup='*')
def ts(a):return scripts('T',sub=a,sup='*')
def ks(a):return scripts('K',sub=a,sup='*')
def ms(a):return scripts('M',sub=a,sup='*')
def deriv(v):return seq(scripts('∂',sub=v),qs())


def equation_rows(i):
    cx=scripts('c',sub='x');cy=scripts('c',sub='y')
    if i==0:
        return [seq('ξ = ',frac(seq('x − ',cx),'D'),'     η = ',frac(seq('y − ',cy),'D'),'     D = 2R')]
    if i==1:
        ps=scripts('ψ',sup='*')
        return [seq(scripts('u',sup='*'),' = ',frac(seq('∂',ps),'∂η'),'     ',scripts('v',sup='*'),' = −',frac(seq('∂',ps),'∂ξ'))]
    if i==2:return [seq(ts('a'),' = ',ks('a'),' + ',ms('a'))]
    if i==3:
        return [seq(ks('R'),' = −',frac('2','D'),paren(seq('ξ',deriv('ξ'),' + η',deriv('η')))),
                seq(ks(cx),' = −',frac('1','D'),deriv('ξ'),'     ',ks(cy),' = −',frac('1','D'),deriv('η'))]
    if i==4:
        return [seq(scripts('T',sub='a'),' = ',scripts('U',sub='0'),paren(seq(ks('a'),' + ',ms('a'))),' + ',scripts('U',sub='a',sup='′'),scripts('q',sub='0',sup='*'))]
    if i==5:
        return [seq(scripts('e',sub='T'),' ≤ ',scripts('κ',sub='K'),scripts('e',sub='K'),' + ',scripts('κ',sub='M'),scripts('e',sub='M'),'     ',scripts('κ',sub='K'),' = ',frac('‖K‖','‖T‖'),'     ',scripts('κ',sub='M'),' = ',frac('‖M‖','‖T‖'))]
    if i==6:
        return [seq('A(ξ,η) = ',frac('ξ − 1.7832861',seq('cos',paren('0.71782035η'),' − 0.3375839')),' − 0.68956524')]
    if i==7:
        return [seq(scripts('ψ',sub='0',sup='*'),'(ξ,η) = η + sin',paren(seq('η − ',frac('η','A(ξ,η)'),' − 0.054244295')))]
    if i==8:
        return [seq('F(ξ,η) = η',paren(seq('1.3596838 − ',scripts('η',sup='2'))),' + cos',paren(seq('ξ + ',scripts(paren('η + 0.38234538'),sup='2'))))]
    if i==9:
        return [seq(scripts('φ',sub='R'),'(ξ,η) = ',paren('11.504055 − ξ'),paren(seq('F(ξ,η) − cos',paren(seq('sin',paren(seq('sin ξ')))),' + 0.19094007')))]
    raise ValueError(i)


def add_equation(doc,i):
    for rr in equation_rows(i):
        p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before=Pt(3);p.paragraph_format.space_after=Pt(7)
        p.paragraph_format.keep_together=True
        om=OxmlElement('m:oMath')
        for r in rr:om.append(r)
        p._p.append(om)


def table_widths(headers):
    n=len(headers)
    if n==2:return [4.1,12.5]
    if n==3:
        if '代表性结果' in headers:return [3.4,8.0,5.2]
        if '优先级' in headers:return [1.4,9.0,6.2]
        return [5.3,5.65,5.65]
    if n==4:
        if headers[0]=='基准梯度来源':return [5.0,3.1,3.8,4.7]
        return [4.4,4.0,4.1,4.1]
    if n==5:return [2.4,3.55,3.55,3.55,3.55]
    return [WIDTH_CM/n]*n


def add_table(doc,rows):
    table=doc.add_table(rows=1,cols=len(rows[0]))
    table.alignment=WD_TABLE_ALIGNMENT.CENTER;table.autofit=False
    widths=table_widths(rows[0])
    for col,w in zip(table.columns,widths):col.width=Cm(w)
    pr=table._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
    pr.append(borders)
    for ri,row in enumerate(rows):
        cells=table.rows[0].cells if ri==0 else table.add_row().cells
        trpr=table.rows[ri]._tr.get_or_add_trPr()
        cant=OxmlElement('w:cantSplit');trpr.append(cant)
        if ri==0:
            h=OxmlElement('w:tblHeader');trpr.append(h)
        for ci,(cell,val,w) in enumerate(zip(cells,row,widths)):
            cell.width=Cm(w);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp=cell._tc.get_or_add_tcPr()
            sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'DCE6ED' if ri==0 else ('F5F7F9' if ri%2==0 else 'FFFFFF'));cp.append(sh)
            margins=OxmlElement('w:tcMar')
            for edge,num in [('top',75),('bottom',75),('left',100),('right',100)]:
                e=OxmlElement('w:'+edge);e.set(qn('w:w'),str(num));e.set(qn('w:type'),'dxa');margins.append(e)
            cp.append(margins)
            p=cell.paragraphs[0];p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.12
            p.paragraph_format.keep_with_next=(ri==0)
            p.alignment=WD_ALIGN_PARAGRAPH.LEFT if ci==0 or (len(val)>22 and ri>0) else WD_ALIGN_PARAGRAPH.CENTER
            inline(p,val)
            for run in p.runs:run.font.size=Pt(9.3);run.bold=(ri==0) or run.bold
    p=doc.add_paragraph();p.paragraph_format.space_after=Pt(1);p.paragraph_format.space_before=Pt(0);p.paragraph_format.line_spacing=1
    p.add_run().font.size=Pt(2)


def build():
    doc=Document();sec=doc.sections[0]
    sec.page_width=Cm(21);sec.page_height=Cm(29.7)
    sec.top_margin=Cm(1.9);sec.bottom_margin=Cm(1.8);sec.left_margin=Cm(2.2);sec.right_margin=Cm(2.2)
    sec.header_distance=Cm(.7);sec.footer_distance=Cm(.8)
    font_style(doc.styles['Normal'],10.5)
    normal=doc.styles['Normal'].paragraph_format
    normal.line_spacing=1.16;normal.space_after=Pt(5);normal.widow_control=True
    font_style(doc.styles['Title'],23,True,'微软雅黑')
    doc.styles['Title'].paragraph_format.space_after=Pt(12)
    for name,size in [('Heading 1',15),('Heading 2',12)]:
        st=doc.styles[name];font_style(st,size,True,'微软雅黑')
        st.paragraph_format.space_before=Pt(12);st.paragraph_format.space_after=Pt(7)
        st.paragraph_format.keep_with_next=True
    font_style(doc.styles['Caption'],9,False)
    doc.styles['Caption'].paragraph_format.space_after=Pt(9)
    doc.styles['Caption'].paragraph_format.keep_together=True
    doc.styles['Caption'].font.italic=False
    # The bundled base document may carry Word's decorative Title border.
    # Remove inherited and direct paragraph rules, leaving only table borders.
    for root in [doc.styles.element, doc._element]:
        for border in list(root.xpath('.//w:pBdr')):
            border.getparent().remove(border)
    # Page numbers aid navigation in a substantial research report.
    p=sec.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run('圆柱绕流研究成果汇总    ');r.font.size=Pt(8)
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');p._p.append(fld)
    props=doc.core_properties;props.title='圆柱绕流符号建模与灵敏度研究成果汇总报告'
    props.subject='项目成果与验证边界';props.author='';props.keywords='PISR;流函数;灵敏度;符号回归;成果汇总'
    lines=MD.read_text(encoding='utf-8').splitlines();i=0;eq=0;figures=0;tables=0
    while i<len(lines):
        line=lines[i].strip()
        if not line:i+=1;continue
        if line=='$$':
            end=i+1
            while end<len(lines) and lines[end].strip()!='$$':end+=1
            assert end<len(lines),'Unclosed equation'
            add_equation(doc,eq);eq+=1;i=end+1;continue
        if line.startswith('|'):
            rr=[]
            while i<len(lines) and lines[i].strip().startswith('|'):
                cols=[v.strip() for v in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r':?-+:?',c) for c in cols):rr.append(cols)
                i+=1
            add_table(doc,rr);tables+=1;continue
        img=re.fullmatch(r'!\[([^]]+)\]\(([^)]+)\)',line)
        if img:
            path=MD.parent/img.group(2)
            p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next=True;p.paragraph_format.space_before=Pt(4);p.paragraph_format.space_after=Pt(3)
            shape=p.add_run().add_picture(str(path),width=Cm(WIDTH_CM))
            shape._inline.docPr.set('descr',img.group(1))
            cp=doc.add_paragraph(img.group(1),style='Caption');cp.alignment=WD_ALIGN_PARAGRAPH.CENTER
            figures+=1;i+=1;continue
        if line.startswith('# '):
            p=doc.add_paragraph(style='Title');inline(p,line[2:]);i+=1;continue
        if line.startswith('## '):
            heading=line[3:]
            p=doc.add_paragraph(heading,style='Heading 1')
            if heading.startswith('1 研究') or heading.startswith('附录A'):
                p.paragraph_format.page_break_before=True
            i+=1;continue
        if line.startswith('### '):
            text=re.sub(r'^(\d+)\.(\d+) ',r'\1 \2 ',line[4:])
            doc.add_paragraph(text,style='Heading 2');i+=1;continue
        if line.startswith('- '):
            p=doc.add_paragraph();p.paragraph_format.space_after=Pt(5)
            p.paragraph_format.keep_together=True
            inline(p,line[2:])
            for r in p.runs:r.font.size=Pt(9.5)
            i+=1;continue
        p=doc.add_paragraph()
        if line.startswith('截至 '):
            p.paragraph_format.space_after=Pt(12)
            p.add_run(line).font.size=Pt(11)
        else:
            inline(p,line)
            p.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
        i+=1
    assert eq==10,eq
    assert figures==6,figures
    doc.save(OUT)
    meta={'output':str(OUT),'native_equation_groups':eq,'figures':figures,'tables':tables,'paragraphs':len(doc.paragraphs),'source':str(MD)}
    (ASSETS/'document_build.json').write_text(json.dumps(meta,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(meta,ensure_ascii=False))

if __name__=='__main__':build()
